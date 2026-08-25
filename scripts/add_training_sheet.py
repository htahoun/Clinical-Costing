import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

P='/home/user/Clinical-Costing/output/Clinical_Costing_2025_MF4159.xlsx'
wb=openpyxl.load_workbook(P)
if 'From 4 codes to 1 claim' in wb.sheetnames:
    del wb['From 4 codes to 1 claim']
ws=wb.create_sheet('From 4 codes to 1 claim', 1)  # right after Sheet1

INK='FF14232A'; INK2='FF405058'; TEAL='FF0D6B6B'; OCHRE='FF8A6410'
TEAL_F=PatternFill('solid',fgColor='FFE4EFEE'); GREY_F=PatternFill('solid',fgColor='FFEFF3F3')
SUM_F=PatternFill('solid',fgColor='FFFFF4D6'); XML_F=PatternFill('solid',fgColor='FFF7F9F9')
thin=Side(style='thin',color='FFDCE4E4')
BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
A=lambda sz=10,b=False,c=INK,i=False: Font(name='Arial',size=sz,bold=b,color=c,italic=i)
M=lambda sz=10,b=False,c=INK: Font(name='Consolas',size=sz,bold=b,color=c)

def put(cell,val,font=None,fill=None,fmt=None,align=None,border=False):
    ws[cell]=val
    if font: ws[cell].font=font
    if fill: ws[cell].fill=fill
    if fmt: ws[cell].number_format=fmt
    if align: ws[cell].alignment=align
    if border: ws[cell].border=BOX

widths={'A':3,'B':10,'C':30,'D':12,'E':12,'F':10,'G':11,'H':10,'I':10,'J':11,'K':11,'L':13,'M':9,'N':9,'O':11,'P':60}
for col,w in widths.items(): ws.column_dimensions[col].width=w

# ---------- title ----------
put('B2','FROM 4 CODES TO 1 CLAIM',A(16,True))
put('B3','How the treatment codes you enter become one cost block in the DOH XML — a real patient from our own 2025 data.',A(10,c=INK2))
put('B4','Training page for the coding / MOI team. Nothing on this sheet needs filling in — it is a worked example.',A(9,c=INK2,i=True))

# ---------- the patient ----------
put('B6','THE VISIT',A(11,True,TEAL))
meta=[('Claim / Encounter ID','728MF4159-c5635e'),('Patient file','35205'),
      ('Dentist','GD5633  ·  General Dentist'),('Date','04/01/2025, 17:33 – 17:53'),
      ('Principal diagnosis','K05.322'),('Facility','MF4159  ·  Vision Dental Clinic')]
r=7
for k,v in meta:
    put(f'B{r}',k,A(9,True,c=INK2)); ws.merge_cells(f'B{r}:D{r}')
    put(f'E{r}',v,A(9)); ws.merge_cells(f'E{r}:H{r}')
    r+=1

# ---------- step 1: the codes ----------
put('B14','STEP 1 — WHAT THE CODERS ENTERED: four treatment codes, four rows in the Excel',A(11,True,TEAL))
hdrs=['Code','Treatment','Anaesthesia','Imaging','OP','Supplies','SWNurs','SWDoc','SWAllied','SWNonClin','Depreciation','OHF','OHC','Row total']
cols=list('BCDEFGHIJKLMNO')
hr=15
for c,h in zip(cols,hdrs):
    put(f'{c}{hr}',h,A(9,True),GREY_F,border=True)
ROWTOTS={16:74.179,17:108.261,18:72.356,19:58.073}
COLTOTS=[18.591, 35.103, 15.006, 18.406, 45.46, 85.359, 25.544, 19.133, 31.727, 10.489, 8.051]
CLAIMTOT=312.869
data=[
 ('01204','Examination, specific',0,0,7.309,3.687,14.998,29.710,7.772,7.374,0.207,1.503,1.619),
 ('02601','Panoramic X-ray',0,35.103,5.058,5.104,5.190,15.421,10.757,5.103,19.078,4.523,2.924),
 ('11111','Scaling, one unit',10.328,0,1.466,3.698,15.043,22.349,3.897,3.698,6.912,3.016,1.949),
 ('11101','Polishing, one unit',8.263,0,1.173,5.917,10.229,17.879,3.118,2.958,5.530,1.447,1.559),
]
r=16
for row in data:
    put(f'B{r}',row[0],M(9),border=True)
    put(f'C{r}',row[1],A(9),border=True)
    for c,v in zip('DEFGHIJKLMN',row[2:]):
        put(f'{c}{r}',v,A(9),fmt='0.000',border=True)
    put(f'O{r}',ROWTOTS[r],A(9,True),fmt='0.000',border=True)
    r+=1

# ---------- step 2: the sum ----------
put('B21','STEP 2 — WHAT THE SCRIPT DOES: add the four rows together. The codes disappear; only the totals remain.',A(11,True,TEAL))
hr=22
for c,h in zip(cols,hdrs):
    put(f'{c}{hr}',h if c not in 'BC' else ('' if c=='B' else 'ONE claim — all four codes summed'),A(9,True),SUM_F,border=True)
r=23
put(f'B{r}','—',A(9,True,OCHRE),SUM_F,border=True)
put(f'C{r}','Claim 728MF4159-c5635e',A(9,True),SUM_F,border=True)
for c,v in zip('DEFGHIJKLMN',COLTOTS):
    put(f'{c}{r}',v,A(9,True,OCHRE),SUM_F,fmt='0.000',border=True)
put(f'O{r}',CLAIMTOT,A(10,True,OCHRE),SUM_F,fmt='0.000',border=True)
put('B25','Check: the claim total (312.869) equals the four row totals added together (74.179 + 108.261 + 72.356 + 58.073). Same money, regrouped — nothing lost, nothing invented.',A(9,i=True,c=INK2))

# ---------- step 3: XML ----------
put('B27','STEP 3 — WHAT DOH RECEIVES: that one summed row, written as one XML block. Amounts below are the Step 2 totals, written into the XML.',A(11,True,TEAL))
xml_rows=[
 ('<Claim>',None),
 ('  <ID>728MF4159-c5635e</ID>',None),
 ('  <ProviderID>MF4159</ProviderID>',None),
 ('  <Encounter>',None),
 ('    <FacilityID>MF4159</FacilityID>  <PatientID>35205</PatientID>  <Type>1</Type>',None),
 ('    <ActivityClinicianID>GD5633</ActivityClinicianID>',None),
 ('    <Start>04/01/2025 17:33</Start>  <End>04/01/2025 17:53</End>',None),
 ('    <Diagnosis><Type>Principal</Type><Code>K05.322</Code></Diagnosis>',None),
 ('    <Procedure><Type/><Code/><Duration>0</Duration></Procedure>   (empty: no theatre)',None),
 ('    <CostBucketDirect>',None),
 ('      <Anaesthesia>','D23'),
 ('      <Imaging>','E23'),
 ('      <OP>','F23'),
 ('      <Supplies>','G23'),
 ('      ...all other buckets 0.000...',None),
 ('    </CostBucketDirect>',None),
 ('    <CostBucketOverheads>  ...all 0.000 for our clinic...  </CostBucketOverheads>',None),
 ('    <CostTypeDetails>',None),
 ('      <SWNurs>','H23'),
 ('      <SWDoc>','I23'),
 ('      <SWAllied>','J23'),
 ('      <SWNonClin>','K23'),
 ('      <Depreciation>','L23'),
 ('      <OHF>','M23'),
 ('      <OHC>','N23'),
 ('      ...all other types 0.000...',None),
 ('    </CostTypeDetails>',None),
 ('  </Encounter>',None),
 ('</Claim>',None),
]
r=28
for text,ref in xml_rows:
    put(f'B{r}',text,M(9,c=INK2),XML_F)
    ws.merge_cells(f'B{r}:H{r}')
    for cc in 'CDEFGH': ws[f'{cc}{r}'].fill=XML_F
    if ref:
        col=ref[0]; put(f'I{r}',COLTOTS['DEFGHIJKLMN'.index(col)],M(9,True,TEAL),XML_F,fmt='0.000')
        put(f'J{r}','← from Step 2',A(8,i=True,c=INK2),XML_F)
    r+=1

# ---------- the rules ----------
r+=1
put(f'B{r}','THE FIVE RULES TO REMEMBER',A(11,True,TEAL)); r+=1
rules=[
 '1.  One <Claim> block per claim — NEVER one per code. 23,983 code rows become 9,718 claim blocks.',
 '2.  The CPT codes themselves do NOT appear in the XML. They are only used to calculate the cost split.',
 '3.  One diagnosis only — the Principal ICD-10 code. DOH rejects files with secondary diagnoses added.',
 '4.  Procedure code stays EMPTY unless the patient went to theatre (TheatreAttendanceFlag = 1).',
 '5.  Nobody types XML by hand. The script builds it from this workbook; your job is a correct spreadsheet:',
]
for t in rules:
    put(f'B{r}',t,A(10)); r+=1
put(f'B{r}','     right Claim ID, right dates, right dentist licence, right principal diagnosis — on every row.',A(10,True)); r+=1
r+=1
put(f'B{r}','Reference: DOH Cost Data Collection Technical Document v1.12 and the Production Connecting-to-Shafafiya Guide (docs/shafafiya/).',A(8,i=True,c=INK2))

ws.sheet_view.showGridLines=False
wb.save(P)
print("sheet added")
