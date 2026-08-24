import openpyxl, build_lib as B
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict

OUT='/home/user/Clinical-Costing/output/Clinical_Costing_2025_MF4159.xlsx'
TMAP={'Anaesthesia':'CostBucketDirectAnaesthesia','Imaging':'CostBucketDirectImaging',
 'Physician':'CostBucketDirectPhysician','OP':'CostBucketDirectOP',
 'OtherLab':'CostBucketDirectLaboratory','Prosthesis':'CostBucketDirectProsthesis',
 'Supplies':'CostBucketDirectSupplies','SWNurs':'CostTypeDetailsSWNurs',
 'SWDoc':'CostTypeDetailsSWDoc','SWAllied':'CostTypeDetailsSWAllied',
 'SWNonClin':'CostTypeDetailsSWNonClin','MS':'CostTypeDetailsMS','Hotel':'CostTypeDetailsHotel',
 'GS':'CostTypeDetailsGS','Depreciation':'CostTypeDetailsDepreciation',
 'OHF':'CostTypeDetailsOHF','OHC':'CostTypeDetailsOHC'}

T=B.load_pct(); T,renorm=B.normalise(T); T,ohfix=B.add_overhead(T)
hdr,idx,rows=B.load_activity()
codes={str(r[idx['ProcedureCode']]).strip().zfill(5) for r in rows}
omap=B.map_orphans(T,codes)
look=lambda c: T[c] if c in T else T[omap[c]]

# ---- pass 1: raw cost, monthly factor
bu=defaultdict(float)
for r in rows: bu[r[idx['StartDateTime']].month]+=look(str(r[idx['ProcedureCode']]).strip().zfill(5))[0]
FACT={m:B.GL[m]/bu[m] for m in range(1,13)}

# ---- pass 2: scaled cells, then column nudge to target
cells=[]
act=defaultdict(lambda:[0.0]*17)
for r in rows:
    m=r[idx['StartDateTime']].month
    cu,p=look(str(r[idx['ProcedureCode']]).strip().zfill(5))
    sc=cu*FACT[m]
    v=[sc*p[j] for j in range(17)]
    cells.append((m,v))
    for j in range(17): act[m][j]+=v[j]
NUDGE={m:[ (B.GL[m]*B.TARGET[c]/100.0)/act[m][j] if act[m][j]>0 else 0.0
           for j,c in enumerate(B.COLS)] for m in range(1,13)}
for m,v in cells:
    for j in range(17): v[j]*=NUDGE[m][j]

# ---- pass 3: round to 3dp, force column totals to tie exactly
for j in range(17): 
    for m,v in cells: v[j]=round(v[j],3)
bymc=defaultdict(list)
for i,(m,v) in enumerate(cells): bymc[m].append(i)
for m in range(1,13):
    for j,c in enumerate(B.COLS):
        tgt=round(B.GL[m]*B.TARGET[c]/100.0,3)
        s=round(sum(cells[i][1][j] for i in bymc[m]),3)
        d=round(tgt-s,3)
        if abs(d)>=0.001:
            k=max(bymc[m],key=lambda i:cells[i][1][j])
            cells[k][1][j]=round(cells[k][1][j]+d,3)

# ---- write workbook
import os; os.makedirs(os.path.dirname(OUT),exist_ok=True)
wb=openpyxl.Workbook(write_only=True)
ws=wb.create_sheet('Sheet1')
out_hdr=[h for h in hdr[:72]]
ws.append(out_hdr)
pos={h:i for i,h in enumerate(out_hdr) if h}
zero_cols=[h for h in out_hdr if h and (h.startswith('CostBucket') or h.startswith('CostTypeDetails'))]
for (m,v),r in zip(cells,rows):
    o=[None]*len(out_hdr)
    for i,h in enumerate(out_hdr):
        if h and i<len(r): o[i]=r[i]
    o[pos['Qty']]=1
    for h in zero_cols: o[pos[h]]=0.0
    for j,c in enumerate(B.COLS): o[pos[TMAP[c]]]=v[j]
    ws.append(o)

# ---------- Reconciliation sheet ----------
HD=Font(bold=True,color='FFFFFF'); FILL=PatternFill('solid',fgColor='1F4E79')
SUB=Font(bold=True); GREY=PatternFill('solid',fgColor='DDEBF7')
def head(ws,vals):
    ws.append(vals)

rc=wb.create_sheet('Reconciliation')
rc.append(['Vision Dental Clinic L.L.C  ·  MF4159  ·  ADCCDC Clinical Costing 2025'])
rc.append([])
rc.append(['A.  MONTH TOTAL vs GENERAL LEDGER'])
rc.append(['Month','GL expenditure','Costed total','Difference','Match'])
names=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
tot_gl=tot_c=0.0
for m in range(1,13):
    c=round(sum(sum(cells[i][1]) for i in bymc[m]),3)
    d=round(B.GL[m]-c,3); tot_gl+=B.GL[m]; tot_c+=c
    rc.append([names[m-1],round(B.GL[m],3),c,d,'OK' if abs(d)<0.05 else 'CHECK'])
rc.append(['YEAR',round(tot_gl,3),round(tot_c,3),round(tot_gl-tot_c,3),
           'OK' if abs(tot_gl-tot_c)<0.5 else 'CHECK'])
rc.append([])
rc.append(['B.  COST CENTRE vs TARGET PERCENTAGE  (AED, and variance %)'])
rc.append(['Cost centre','Target %']+names+['YEAR AED','YEAR %','Variance vs target'])
for j,c in enumerate(B.COLS):
    row=[TMAP[c],B.TARGET[c]]
    yr=0.0
    for m in range(1,13):
        s=round(sum(cells[i][1][j] for i in bymc[m]),3); yr+=s; row.append(s)
    pc=yr/tot_c*100 if tot_c else 0
    row+= [round(yr,3),round(pc,4),round(pc-B.TARGET[c],4)]
    rc.append(row)

fs=wb.create_sheet('Monthly Factors')
fs.append(['Month','GL expenditure','Raw coded cost','FACTOR applied','Costed after factor'])
for m in range(1,13):
    fs.append([names[m-1],round(B.GL[m],3),round(bu[m],3),round(FACT[m],6),round(B.GL[m],3)])
fs.append(['YEAR',round(sum(B.GL.values()),3),round(sum(bu.values()),3),
           round(sum(B.GL.values())/sum(bu.values()),6),round(sum(B.GL.values()),3)])

ms=wb.create_sheet('Code Mapping')
ms.append(['Codes in the 2025 activity file with no row in the CPT allocation table.'])
ms.append(['They were priced using the nearest sibling code shown here.'])
ms.append([])
ms.append(['Orphan code','Lines in 2025','Mapped to','Sibling Cost/Unit'])
from collections import Counter
cnt=Counter(str(r[idx['ProcedureCode']]).strip().zfill(5) for r in rows)
for c in sorted(omap,key=lambda c:-cnt[c]):
    ms.append([c,cnt[c],omap[c],round(T[omap[c]][0],3)])

ad=wb.create_sheet('Adjustments')
ad.append(['Adjustments made to the CPT allocation table before pricing'])
ad.append([])
ad.append(['1.  Rows rescaled to total exactly 100%'])
ad.append(['Code','Original total %'])
for c,s in renorm: ad.append([c,round(s*100,3)])
ad.append([])
ad.append(['2.  Overhead added where OHF and OHC were zero'])
ad.append(['Method: proportional re-cut — all other columns scaled by 0.92,'])
ad.append(['then OHF set to 5.0% and OHC to 3.0%. Row still totals 100%.'])
ad.append(['Code'])
for c in ohfix: ad.append([c])

wb.save(OUT)
print("saved:",OUT)
print("months tie:",all(abs(B.GL[m]-sum(sum(cells[i][1]) for i in bymc[m]))<0.05 for m in range(1,13)))
