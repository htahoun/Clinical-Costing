import openpyxl, build2_lib as B
from collections import defaultdict

OUT='/home/user/Clinical-Costing/output/Clinical_Costing_2025_MF4159_v2.xlsx'
TMAP={'Anaesthesia':'CostBucketDirectAnaesthesia','Imaging':'CostBucketDirectImaging',
 'Physician':'CostBucketDirectPhysician','OP':'CostBucketDirectOP',
 'OtherLab':'CostBucketDirectLaboratory','Prosthesis':'CostBucketDirectProsthesis',
 'Supplies':'CostBucketDirectSupplies','SWNurs':'CostTypeDetailsSWNurs',
 'SWDoc':'CostTypeDetailsSWDoc','SWAllied':'CostTypeDetailsSWAllied',
 'SWNonClin':'CostTypeDetailsSWNonClin','MS':'CostTypeDetailsMS','Hotel':'CostTypeDetailsHotel',
 'GS':'CostTypeDetailsGS','Depreciation':'CostTypeDetailsDepreciation',
 'OHF':'CostTypeDetailsOHF','OHC':'CostTypeDetailsOHC'}

T,fixes=B.load_alloc()
months,hdr0=B.load_activity()

# pass 1: factors
FACT={}; MODEL={}
for mi,(rows,idx) in months.items():
    s=sum(B.lookup(T,r[idx['ProcedureCode']])[0] for r in rows)
    MODEL[mi]=s; FACT[mi]=B.GL[mi]/s

# pass 2: scaled unrounded cells
cells={}   # mi -> list of [17 floats] aligned with rows
for mi,(rows,idx) in months.items():
    L=[]
    for r in rows:
        cu,p=B.lookup(T,r[idx['ProcedureCode']])
        sc=cu*FACT[mi]
        L.append([sc*p[j] for j in range(17)])
    cells[mi]=L

# pass 3: round 3dp, force per-column ties, then month total == round(GL,3)
for mi in range(1,13):
    L=cells[mi]
    colsum=[sum(v[j] for v in L) for j in range(17)]
    tgt=[round(c,3) for c in colsum]
    d=round(round(B.GL[mi],3)-sum(tgt),3)
    jmax=max(range(17),key=lambda j:tgt[j])
    tgt[jmax]=round(tgt[jmax]+d,3)
    for j in range(17):
        for v in L: v[j]=round(v[j],3)
        s=round(sum(v[j] for v in L),3)
        dd=round(tgt[j]-s,3)
        if abs(dd)>=0.001:
            k=max(range(len(L)),key=lambda i:L[i][j])
            L[k][j]=round(L[k][j]+dd,3)

# write workbook
wb=openpyxl.Workbook(write_only=True)
ws=wb.create_sheet('Sheet1')
out_hdr=[h for h in hdr0]
ws.append(out_hdr)
pos={h:i for i,h in enumerate(out_hdr) if h}
costcols=[h for h in out_hdr if h and (h.startswith('CostBucket') or h.startswith('CostTypeDetails'))]
for mi in range(1,13):
    rows,idx=months[mi]
    for r,v in zip(rows,cells[mi]):
        o=[None]*len(out_hdr)
        for h,i in pos.items():
            src=idx.get(h)
            if src is not None and src<len(r): o[i]=r[src]
        for h in costcols: o[pos[h]]=0.0
        tot=0.0
        for j,c in enumerate(B.COLS):
            o[pos[TMAP[c]]]=v[j]; tot+=v[j]
        if 'total cost' in pos: o[pos['total cost']]=round(tot,3)
        ws.append(o)

names=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
rc=wb.create_sheet('Reconciliation')
rc.append(['Vision Dental Clinic L.L.C · MF4159 · ADCCDC Clinical Costing FY2025 · v2'])
rc.append(['Method: scale-only. Each month scaled by one factor so its total ties exactly to the expenditure file.'])
rc.append(['The cost-centre mix is whatever the per-code allocation produces (no golden-percentage nudge).'])
rc.append([])
rc.append(['A. MONTH TOTAL vs EXPENDITURE'])
rc.append(['Month','Expenditure (AED)','Costed total','Difference'])
tg=tc=0.0
for mi in range(1,13):
    c=round(sum(sum(v) for v in cells[mi]),3); d=round(B.GL[mi]-c,3)
    tg+=B.GL[mi]; tc+=c
    rc.append([names[mi-1],round(B.GL[mi],3),c,d])
rc.append(['YEAR',round(tg,3),round(tc,3),round(tg-tc,3)])
rc.append([])
rc.append(['B. RESULTING COST-CENTRE MIX (% of month) — produced by the per-code splits'])
rc.append(['Cost centre']+names+['YEAR %','YEAR AED'])
for j,c in enumerate(B.COLS):
    row=[TMAP[c]]
    yr=0.0
    for mi in range(1,13):
        s=sum(v[j] for v in cells[mi]); yr+=s
        row.append(round(s/B.GL[mi]*100,2))
    row+=[round(yr/tc*100,2),round(yr,3)]
    rc.append(row)

fs=wb.create_sheet('Monthly Factors')
fs.append(['Month','Expenditure (file governs)','Modelled cost (units x price)','FACTOR applied'])
for mi in range(1,13):
    fs.append([names[mi-1],round(B.GL[mi],3),round(MODEL[mi],3),round(FACT[mi],6)])
fs.append(['YEAR',round(sum(B.GL.values()),3),round(sum(MODEL.values()),3),round(sum(B.GL.values())/sum(MODEL.values()),6)])

ad=wb.create_sheet('Code Mapping & Adjustments')
ad.append(['Adjustments applied to the allocation table before pricing (source: CPT_Allocation_Cheiri_updated.xlsx)'])
for f in fixes: ad.append([f])
ad.append([])
ad.append(['Remaining orphan codes priced from their sibling:'])
for k,v in B.ORPHAN_MAP.items(): ad.append([f'{k} -> {v}'])
ad.append([])
ad.append(['Code 99555 (+E Additional Expense of Materials) is priced 0.00 — its lines cost zero by design.'])

wb.save(OUT)
print("saved",OUT)
