"""
Shared library: constants, allocation-table loader with the locked fixes,
activity reader with the locked hygiene rules, exact-fils rounding helpers.

Locked decisions are documented in docs/CLINICAL_COSTING.md §3–§4. Do not change
them here without the Finance Director's instruction.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

FACILITY = "MF4159"

# The 17 allocation columns, in the allocation-table order.
ALLOC_COLS = ["Anaesthesia", "Imaging", "Physician", "OP", "OtherLab", "Prosthesis",
              "Supplies", "SWNurs", "SWDoc", "SWAllied", "SWNonClin", "MS", "Hotel",
              "GS", "Depreciation", "OHF", "OHC"]

# ---- Locked allocation-table fixes (§3) ---------------------------------------
REKEY = {("14611", Decimal("378.00")): "14612"}       # duplicate 14611 @ 378.00 is really 14612
DEDUPE_IDENTICAL = {"12701", "29301"}                 # identical duplicate rows
RESCALE_TO_100 = {"01901", "02504", "66301", "72111", "72119"}
ORPHAN_MAP = {"12702": "12701", "29302": "29301", "42321": "42311", "27113": "27111"}
ZERO_COST_CODES = {"99555", "00000"}                   # +E materials; blank code

# ---- Locked hygiene rules (§4) ------------------------------------------------
KNOWN_MODAL_SPECIALTY = {
    "GD11525": "Oral Surgeon",
    "GD5633": "General Dentist",
    "GD22503": "General Dentist",
    "GD45264": "General Dentist",
    "GD36537": "General Dentist",
}

# Canonical activity columns and the header spellings accepted for each.
ACTIVITY_COLS = {
    "claim_id":   ["Claim ID", "ClaimID", "Claim.ID", "Claim Ref", "Claim Reference", "Claim No", "ClaimNo", "Claim"],
    "serial":     ["Serial", "Activity Serial", "Serial No", "Serial No.", "Encounter Serial", "Activity ID",
                   "ActivityID", "Encounter ID", "Encounter.ID"],
    "patient_id": ["FileNo.", "FileNo", "File No", "File No.", "Patient ID", "PatientID", "MRN"],
    "clinician":  ["Doctor No", "Doctor No.", "DoctorNo", "Clinician", "Clinician ID", "ActivityClinicianID",
                   "Doctor"],
    "specialty":  ["Speciality", "Specialty"],
    "start":      ["StartDateTime", "Start Date Time", "Start DateTime", "Encounter Start", "Start", "Date"],
    "end":        ["EndDateTime", "End Date Time", "End DateTime", "Encounter End", "End"],
    "code":       ["Code", "CPT", "CPT Code", "Activity Code", "ActivityCode", "Service Code", "Procedure Code"],
    "units":      ["Units", "Quantity", "Qty"],
    "diagnosis":  ["Diagnosis", "Principal Diagnosis", "Diagnosis Code", "ICD", "ICD Code", "PrincipalDiagnosis",
                   "DiagnosisCode"],
    "theatre":    ["TheatreAttendanceFlag", "Theatre Attendance Flag", "Theatre"],
    "procedure":  ["Procedure", "ProcedureCode", "Theatre Procedure"],
    "duration":   ["Duration", "Procedure Duration"],
    "enc_type":   ["Type", "Encounter Type", "EncounterType"],
    "payer":      ["Payer", "Payer ID", "PayerID"],
    "net":        ["Net", "Net Amount", "Claim Net", "Amount"],
    "source":     ["Source"],
}
REQUIRED_ACTIVITY = ["claim_id", "serial", "patient_id", "clinician", "start", "code"]


def norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def match_headers(headers: list, spec: dict) -> dict:
    """Return {canonical: column index} using tolerant header matching."""
    normed = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    found = {}
    for canon, aliases in spec.items():
        for a in aliases:
            i = normed.get(norm_header(a))
            if i is not None:
                found[canon] = i
                break
    return found


def D(x) -> Decimal:
    if x is None or x == "":
        return Decimal(0)
    if isinstance(x, Decimal):
        return x
    if isinstance(x, float):
        return Decimal(repr(x))
    return Decimal(str(x).replace(",", "").strip())


def to_fils(x) -> int:
    """AED (any numeric) -> integer fils, half-up at 3 dp."""
    return int((D(x) * 1000).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def from_fils(f: int) -> Decimal:
    return Decimal(f) / 1000


def clean_code(c) -> str:
    if c is None:
        return "00000"
    s = str(c).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s or "00000"


def clean_serial(s) -> str:
    """Strip a trailing B (locked rule) and surrounding whitespace."""
    s = str(s or "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s[:-1] if s.upper().endswith("B") else s


def parse_dt(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d-%m-%Y %H:%M", "%d-%m-%Y", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"unparseable date: {v!r}")


# ---- Allocation table ----------------------------------------------------------
def load_alloc(path: str | Path, sheet: str | None = None):
    """
    Read the CPT allocation table. Returns (alloc, log) where
      alloc[code] = {"cost": Decimal cost/unit, "split": {col: Decimal fraction}}
    with the locked fixes applied and logged.
    Expected headers: a code column, a Cost/Unit column, and the 17 ALLOC_COLS
    (percentages summing to 100; fractions summing to 1 are also accepted).
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    spec = {"code": ["Code", "CPT", "CPT Code", "Activity Code", "ActivityCode"],
            "cost": ["Cost/Unit", "Cost per Unit", "CostPerUnit", "Unit Cost", "Cost"],
            **{c: [c] for c in ALLOC_COLS}}
    idx = match_headers(headers, spec)
    missing = [k for k in ("code", "cost", *ALLOC_COLS) if k not in idx]
    if missing:
        raise SystemExit(f"Allocation table {path}: missing columns {missing}; headers were {headers}")

    log: list[str] = []
    raw: list[tuple[str, Decimal, dict]] = []
    for r in rows:
        if r is None or all(v in (None, "") for v in r):
            continue
        code = clean_code(r[idx["code"]])
        cost = D(r[idx["cost"]])
        pct = {c: D(r[idx[c]]) for c in ALLOC_COLS}
        raw.append((code, cost, pct))

    alloc: dict[str, dict] = {}
    seen: dict[str, tuple] = {}
    for code, cost, pct in raw:
        key = (code, cost.quantize(Decimal("0.01")))
        if key in REKEY:
            log.append(f"re-key {code} @ {cost} -> {REKEY[key]} (locked)")
            code = REKEY[key]
        if code in alloc:
            same = (alloc[code]["_cost"] == cost and alloc[code]["_pct"] == pct)
            if code in DEDUPE_IDENTICAL and same:
                log.append(f"dedupe identical row {code} (locked)")
                continue
            log.append(f"WARNING duplicate code {code} with {'identical' if same else 'DIFFERENT'} values; "
                       f"first row kept — raise with the Finance Director")
            continue
        total = sum(pct.values())
        scale = Decimal(1)
        if total == 0:
            log.append(f"WARNING {code}: all-zero split (no cost split possible)")
        elif abs(total - 100) <= Decimal("0.05"):
            scale = Decimal(100) / total
        elif abs(total - 1) <= Decimal("0.0005"):
            scale = Decimal(1) / total
        else:
            if code in RESCALE_TO_100:
                scale = Decimal(100) / total
                log.append(f"rescale {code} from {total.quantize(Decimal('0.01'))}% to 100% (locked)")
            else:
                scale = Decimal(100) / total
                log.append(f"WARNING {code}: split sums to {total.quantize(Decimal('0.01'))}%, rescaled to 100% — confirm with Cherie")
        frac = {c: (pct[c] * scale / 100 if total != 0 else Decimal(0)) for c in ALLOC_COLS}
        if total != 0 and abs(total - 1) <= Decimal("0.0005"):
            frac = {c: pct[c] * scale for c in ALLOC_COLS}
        if frac["OHF"] == 0 and frac["OHC"] == 0 and total != 0:
            log.append(f"WARNING {code}: zero overheads — every code must carry its overhead share")
        alloc[code] = {"cost": cost, "split": frac, "_cost": cost, "_pct": pct}

    for orphan, parent in ORPHAN_MAP.items():
        if orphan not in alloc and parent in alloc:
            alloc[orphan] = dict(alloc[parent])
            log.append(f"orphan {orphan} priced from sibling {parent} (locked)")
    for z in ZERO_COST_CODES:
        if z not in alloc:
            alloc[z] = {"cost": Decimal(0), "split": {c: Decimal(0) for c in ALLOC_COLS},
                        "_cost": Decimal(0), "_pct": {}}
            log.append(f"code {z} carried at zero cost (locked)")
    for v in alloc.values():
        v.pop("_cost", None)
        v.pop("_pct", None)
    return alloc, log


# ---- Activity --------------------------------------------------------------------
def load_activity(path: str | Path, sheets: list[str] | None = None):
    """
    Read activity lines (one row = one unit of one CPT code inside one claim)
    from one sheet or several per-month tabs. Applies the locked hygiene rules.
    Returns (lines, log). Each line is a dict with canonical keys plus `row_ref`.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = sheets or [ws.title for ws in wb.worksheets]
    lines: list[dict] = []
    log: list[str] = []
    for name in names:
        ws = wb[name]
        rows = ws.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            continue
        idx = match_headers(headers, ACTIVITY_COLS)
        missing = [k for k in REQUIRED_ACTIVITY if k not in idx]
        if missing:
            log.append(f"sheet {name!r} skipped: missing columns {missing}")
            continue
        n = 0
        for rn, r in enumerate(rows, start=2):
            if r is None or all(v in (None, "") for v in r):
                continue
            g = lambda k: (r[idx[k]] if k in idx and idx[k] < len(r) else None)
            start = parse_dt(g("start"))
            if start is None:
                log.append(f"{name}!{rn}: blank StartDateTime — row skipped")
                continue
            units = D(g("units")) if g("units") not in (None, "") else Decimal(1)
            theatre = str(g("theatre") or "0").strip()
            theatre = "1" if theatre in ("1", "1.0", "Y", "y", "True", "TRUE") else "0"
            line = {
                "sheet": name, "row_ref": f"{name}!{rn}",
                "claim_id": str(g("claim_id") or "").strip(),
                "serial": clean_serial(g("serial")),
                "patient_id": str(g("patient_id") or "").strip().removesuffix(".0"),
                "clinician": str(g("clinician") or "").strip(),
                "specialty": str(g("specialty") or "").strip(),
                "start": start,
                "end": parse_dt(g("end")) or start,
                "code": clean_code(g("code")),
                "units": units,
                "diagnosis": str(g("diagnosis") or "").strip(),
                "theatre": theatre,
                "procedure": str(g("procedure") or "").strip(),
                "duration": str(g("duration") or "0").strip().removesuffix(".0") or "0",
                "enc_type": str(g("enc_type") or "1").strip().removesuffix(".0") or "1",
                "payer": str(g("payer") or "").strip(),
                "net": D(g("net")) if g("net") not in (None, "") else None,
                "source": str(g("source") or "").strip(),
            }
            if not line["claim_id"]:
                log.append(f"{line['row_ref']}: blank claim id — row skipped")
                continue
            lines.append(line)
            n += 1
        log.append(f"sheet {name!r}: {n} lines")

    # Blank Speciality -> clinician's modal speciality (locked rule).
    modal: dict[str, Counter] = defaultdict(Counter)
    for ln in lines:
        if ln["specialty"]:
            modal[ln["clinician"]][ln["specialty"]] += 1
    filled = Counter()
    for ln in lines:
        if not ln["specialty"]:
            spec = KNOWN_MODAL_SPECIALTY.get(ln["clinician"]) or \
                (modal[ln["clinician"]].most_common(1)[0][0] if modal[ln["clinician"]] else "")
            if spec:
                ln["specialty"] = spec
                filled[(ln["clinician"], spec)] += 1
            else:
                log.append(f"{ln['row_ref']}: blank speciality and no modal value for {ln['clinician']}")
    for (c, s), n in sorted(filled.items()):
        log.append(f"speciality filled: {c} -> {s} ({n} lines)")
    return lines, log


# ---- Exact-fils rounding --------------------------------------------------------
def force_sum(values: list[int], target: int) -> list[int]:
    """Adjust the largest value so that sum(values) == target (integer fils)."""
    if not values:
        return values
    diff = target - sum(values)
    if diff:
        i = max(range(len(values)), key=lambda k: values[k])
        values[i] += diff
    return values
