#!/usr/bin/env python3
"""
FINAL costing engine — scale-only (locked, docs/CLINICAL_COSTING.md §2).

  python3 scripts/build4_final.py --activity input/activity.xlsx \
      --alloc input/allocation.xlsx --out output/Clinical_Costing_2025_MF4159_FINAL.xlsx

Per month: modelled = Σ Cost/Unit × units; factor = ledger ÷ modelled;
cell = Cost/Unit × units × factor × column %; 3 dp per cell; each column's
rounded sum forced to its rounded target; month total forced to round(GL, 3)
exactly (largest cell absorbs the residual). No golden-percentage nudging.

Writes: Sheet1 (costed lines + Source), Reconciliation, Monthly Factors,
Merge Log. Prints the verification table.
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from build2_lib import ALLOC_COLS, force_sum, from_fils, load_activity, load_alloc  # noqa: E402
from costsub_schema import TMAP  # noqa: E402
from gl_new import gl_fils  # noqa: E402

IDENTITY_COLS = [
    ("Month", "month"), ("Claim ID", "claim_id"), ("Serial", "serial"), ("FileNo.", "patient_id"),
    ("Doctor No", "clinician"), ("Speciality", "specialty"), ("StartDateTime", "start"),
    ("EndDateTime", "end"), ("Code", "code"), ("Units", "units"), ("Diagnosis", "diagnosis"),
    ("TheatreAttendanceFlag", "theatre"), ("Procedure", "procedure"), ("Duration", "duration"),
    ("Type", "enc_type"), ("Payer", "payer"), ("Source", "source"),
]
COST_OUT = [TMAP[c] for c in ALLOC_COLS]   # 17 output columns, Gen1-virtual XML names


def cost_month(lines: list[dict], alloc: dict, gl: int, log: list[str]):
    """Return (cells[line_index][col] in fils, factor Decimal, modelled Decimal)."""
    unit_cost = []
    for ln in lines:
        a = alloc.get(ln["code"])
        if a is None:
            log.append(f"{ln['row_ref']}: code {ln['code']} not in allocation table — costed 0; ADD TO TABLE")
            a = {"cost": Decimal(0), "split": {c: Decimal(0) for c in ALLOC_COLS}}
            alloc[ln["code"]] = a
        unit_cost.append(a["cost"] * ln["units"])
    modelled = sum(unit_cost, Decimal(0))
    if modelled == 0:
        raise SystemExit(f"month {lines[0]['start']:%Y-%m}: modelled cost is zero; cannot scale")
    factor = (Decimal(gl) / 1000) / modelled
    exact = []   # Decimal AED per cell
    for ln, uc in zip(lines, unit_cost):
        split = alloc[ln["code"]]["split"]
        exact.append([uc * factor * split[c] for c in ALLOC_COLS])
    # round every cell to 3 dp (fils)
    cells = [[int((v * 1000).quantize(Decimal(1), rounding=ROUND_HALF_UP)) for v in row] for row in exact]
    # force each column to its rounded target
    for j in range(len(ALLOC_COLS)):
        target = int((sum(row[j] for row in exact) * 1000).quantize(Decimal(1), rounding=ROUND_HALF_UP))
        col = [row[j] for row in cells]
        col = force_sum(col, target)
        for i, v in enumerate(col):
            cells[i][j] = v
    # force the month total to round(GL, 3) exactly: largest cell absorbs
    diff = gl - sum(sum(row) for row in cells)
    if diff:
        i, j = max(((i, j) for i in range(len(cells)) for j in range(len(ALLOC_COLS))),
                   key=lambda ij: cells[ij[0]][ij[1]])
        cells[i][j] += diff
        log.append(f"month {lines[0]['start']:%Y-%m}: residual {diff} fils absorbed by largest cell")
    return cells, factor, modelled


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--activity", required=True, help="accepted activity master workbook")
    ap.add_argument("--alloc", required=True, help="CPT allocation table workbook")
    ap.add_argument("--out", required=True, help="costed workbook to write")
    ap.add_argument("--year", type=int, default=2025)
    ap.add_argument("--sheets", nargs="*", help="activity sheet names (default: all)")
    args = ap.parse_args(argv)

    alloc, alog = load_alloc(args.alloc)
    lines, llog = load_activity(args.activity, args.sheets)
    log = [f"[alloc] {m}" for m in alog] + [f"[activity] {m}" for m in llog]
    if not lines:
        raise SystemExit("no activity lines read")

    by_month: dict[int, list[dict]] = defaultdict(list)
    for ln in lines:
        if ln["start"].year != args.year:
            log.append(f"{ln['row_ref']}: StartDateTime {ln['start']:%Y-%m-%d} outside FY{args.year} — excluded")
            continue
        by_month[ln["start"].month].append(ln)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([h for h, _ in IDENTITY_COLS] + ["Cost/Unit", "Factor"] + COST_OUT + ["total cost"])
    recon, factors = [], []
    print(f"{'Month':<8}{'Lines':>8}{'Claims':>8}{'Modelled':>16}{'Ledger':>16}{'Costed':>16}{'Factor':>10}{'Tie':>6}")
    for m in sorted(by_month):
        mlines = by_month[m]
        gl = gl_fils(args.year, m)
        cells, factor, modelled = cost_month(mlines, alloc, gl, log)
        costed = sum(sum(r) for r in cells)
        claims = len({ln["claim_id"] for ln in mlines})
        for ln, row in zip(mlines, cells):
            ident = []
            for _, k in IDENTITY_COLS:
                v = m if k == "month" else ln[k]
                if k in ("start", "end"):
                    v = v.strftime("%d/%m/%Y %H:%M")
                elif k == "units":
                    v = float(v)
                ident.append(v)
            ws.append(ident + [float(alloc[ln["code"]]["cost"]), float(factor)]
                      + [float(from_fils(f)) for f in row] + [float(from_fils(sum(row)))])
        tie = "OK" if costed == gl else "FAIL"
        recon.append([m, len(mlines), claims, float(modelled), float(from_fils(gl)), float(from_fils(costed)),
                      float(from_fils(costed - gl)), tie])
        factors.append([m, float(factor), float(modelled), float(from_fils(gl))])
        print(f"{args.year}-{m:02d}{len(mlines):>8}{claims:>8}{float(modelled):>16,.3f}"
              f"{float(from_fils(gl)):>16,.3f}{float(from_fils(costed)):>16,.3f}{float(factor):>10.4f}{tie:>6}")

    r = wb.create_sheet("Reconciliation")
    r.append(["Month", "Lines", "Claims", "Modelled cost", "Ledger", "Costed", "Difference", "Tie"])
    for row in recon:
        r.append(row)
    r.append(["TOTAL", sum(x[1] for x in recon), sum(x[2] for x in recon), sum(x[3] for x in recon),
              sum(x[4] for x in recon), sum(x[5] for x in recon), sum(x[6] for x in recon),
              "OK" if all(x[7] == "OK" for x in recon) else "FAIL"])
    f = wb.create_sheet("Monthly Factors")
    f.append(["Month", "Factor", "Modelled cost", "Ledger"])
    for row in factors:
        f.append(row)
    lg = wb.create_sheet("Merge Log")
    lg.append(["Entry"])
    for m in log:
        lg.append([m])
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    bad = [x for x in recon if x[7] != "OK"]
    print(f"\nwrote {args.out}  ({len(log)} log entries; {len(bad)} month(s) failing to tie)")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
