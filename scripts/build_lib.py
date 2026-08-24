import openpyxl
from collections import defaultdict

COLS=['Anaesthesia','Imaging','Physician','OP','OtherLab','Prosthesis','Supplies',
      'SWNurs','SWDoc','SWAllied','SWNonClin','MS','Hotel','GS','Depreciation','OHF','OHC']
CI={c:i for i,c in enumerate(COLS)}
# target % of month expenditure (the 16-line list; Physician = 0)
TARGET={'Anaesthesia':4.8,'Imaging':3.6,'Physician':0.0,'OP':2.1,'OtherLab':7.0,
        'Prosthesis':6.2,'Supplies':18.6,'SWNurs':12.0,'SWDoc':25.0,'SWAllied':4.8,
        'SWNonClin':3.9,'MS':0.4,'Hotel':0.4,'GS':0.1,'Depreciation':7.2,'OHF':2.1,'OHC':1.8}
GL={1:650297.845,2:698511.42,3:595631.0,4:492438.75,5:655920.685,6:603922.605,
    7:836689.59,8:848189.225,9:610963.44,10:672185.07,11:781176.49,12:657586.385}
PCT_XLSX='/root/.claude/uploads/ad1e78a6-69df-510f-b708-f3b8c52824be/272e96c3-CPT_Allocation_Claude.xlsx'
ACT_XLSX='/root/.claude/uploads/ad1e78a6-69df-510f-b708-f3b8c52824be/8986d8ec-Clinical_2025_one_sheet_Claude.xlsx'

def load_pct():
    wb=openpyxl.load_workbook(PCT_XLSX,data_only=True,read_only=True)
    ws=wb['Pecentage Allocation']; T={}
    for i,r in enumerate(ws.iter_rows(values_only=True),1):
        if i==1 or r[0] is None: continue
        T[str(int(r[0])).zfill(5)]=(float(r[2] or 0),[float(r[3+j] or 0) for j in range(17)])
    return T

def normalise(T):
    """Rows must total exactly 1.0. Rescale the few that don't."""
    out={}; fixed=[]
    for c,(cu,p) in T.items():
        s=sum(p)
        if s==0: out[c]=(cu,p); continue
        if abs(s-1.0)>1e-9:
            p=[v/s for v in p]; fixed.append((c,s))
        out[c]=(cu,p)
    return out,fixed

def add_overhead(T, ohf=0.05, ohc=0.03):
    """Proportional re-cut: make room for OHF/OHC on rows that carry none."""
    out={}; touched=[]
    for c,(cu,p) in T.items():
        if sum(p)==0: out[c]=(cu,p); continue
        if p[CI['OHF']]==0 and p[CI['OHC']]==0:
            k=1.0-(ohf+ohc)
            q=[v*k for v in p]; q[CI['OHF']]=ohf; q[CI['OHC']]=ohc
            out[c]=(cu,q); touched.append(c)
        else: out[c]=(cu,p)
    return out,touched

def load_activity():
    wb=openpyxl.load_workbook(ACT_XLSX,data_only=True,read_only=True); ws=wb['Sheet1']
    it=ws.iter_rows(values_only=True); hdr=list(next(it))
    idx={h:i for i,h in enumerate(hdr) if h}
    rows=[r for r in it if not all(v is None or str(v).strip()=='' for v in r)]
    return hdr,idx,rows

def map_orphans(T, codes):
    """Nearest sibling: longest shared prefix, then closest numerically."""
    known=sorted(T)
    m={}
    for c in sorted(codes):
        if c in T: continue
        best=None
        for k in known:
            if k=='99555': continue
            pre=0
            for a,b in zip(c,k):
                if a==b: pre+=1
                else: break
            d=abs(int(c)-int(k))
            key=(-pre,d)
            if best is None or key<best[0]: best=(key,k)
        m[c]=best[1]
    return m
