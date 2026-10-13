"""
Split a month's Cost of Goods Sold into DOH cost types from the clinic's stock
movement sheet (monthly Financial Report, sheet "Movement <Mon> YYYY").

Usage (as a module from fill_ledger_map_from_financials.py, or directly):
  python3 scripts/cogs_from_movement.py input/source/Financial_Report_Jan_2025_1.xlsx

The sheet records every stock transaction with Category / Sub-Category and a
Transaction Value. Consumption = "02- Usage" (negative) net of "03- Internal
Return" (positive; unused implants returned to stock). "04- Exchange Received"
and "05- Exchange Returned" cancel out. External laboratory invoices are not
stock and are added from the P&L line. Verified on January 2025: usage
132,757.49 − returns 32,500.00 + external lab 1,870.00 = 102,127.49 = the
Monthly P&L Cost of Goods Sold.
"""
from __future__ import annotations

import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

D2 = Decimal("0.01")

# (category prefix, sub-category prefix or None) -> DOH cost type. First match wins; order matters.
CATEGORY_MAP = [
    ("Imaging", None, "Imaging"),                                  # any imaging / X-ray / radiograph category
    ("X-Ray", None, "Imaging"),
    ("Radiograph", None, "Imaging"),
    ("02- Drugs", "01- Local Anaesthetics", "Pharmacy"),          # local anaesthetic cartridges
    ("04- Dental Materials", "02- Nitrous & Oxygen", "Pharmacy"),  # sedation gases
    ("02- Drugs", None, "MS"),                                     # needles, syringes, vital-signs and emergency supplies
    ("05- Implants", None, "Prostheses"),                          # fixtures, healing/cover screws, abutments, impression
                                                                   # tools, bone-graft materials (implants and accessories)
    ("06- Laboratory", None, "Prostheses"),                        # blocks, blanks, 3D inks, stone, stains, milling tools:
                                                                   # in-house prosthesis materials (ruling 13 Oct 2026: no laboratory)
    ("04- Dental Materials", None, "MS"),                          # restorative, endodontic, burs, impression, bleaching…
    ("03- Prevention & Control of Infection", None, "MS"),         # PPE, disinfectants, wraps, sterilisation monitoring
    ("01- Medical & LAB Devices", None, "MS"),                     # hand instruments and small devices expensed on issue
    ("08- Facility Management", None, "OHF"),                      # pantry, stationery
]
CONSUMING = {"02- Usage", "03- Internal Return"}


def dec(v) -> Decimal:
    if v in (None, "", "--", "NA") or not isinstance(v, (int, float, Decimal, str)):
        return Decimal(0)
    try:
        return Decimal(str(v).replace(",", ""))          # exact; rounded once per cost type at the end
    except Exception:  # noqa: BLE001
        return Decimal(0)


def cost_type_for(category: str, sub: str) -> str:
    for cat_prefix, sub_prefix, ct in CATEGORY_MAP:
        hit = category.startswith(cat_prefix) or (sub_prefix is None and cat_prefix.lower() in (category + " " + sub).lower())
        if hit and (sub_prefix is None or sub.startswith(sub_prefix)):
            return ct
    raise KeyError(f"no mapping for category {category!r} / {sub!r}")


def cogs_split(report_path: str | Path, external_lab: Decimal = Decimal(0)) -> tuple[dict[str, Decimal], list[str]]:
    """Return ({cost type: AED consumed}, log). Consumption is positive."""
    wb = openpyxl.load_workbook(report_path, data_only=True, read_only=True)
    ws = next(s for s in wb.worksheets if s.title.lower().startswith("movement"))
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
    hi = {h: i for i, h in enumerate(hdr)}
    col_t, col_c, col_s, col_v = hi["Transaction Type"], hi["Category"], hi["Sub-Category"], hi["Transaction Value"]
    out: dict[str, Decimal] = defaultdict(Decimal)
    detail: dict[tuple, Decimal] = defaultdict(Decimal)
    for r in rows:
        if not r or r[col_t] not in CONSUMING:
            continue
        cat, sub = str(r[col_c]).strip(), str(r[col_s]).strip()
        v = -dec(r[col_v])                                   # usage is negative in the sheet; returns positive
        ct = cost_type_for(cat, sub)
        out[ct] += v
        detail[(ct, cat, sub)] += v
    out = {ct: v.quantize(D2, rounding=ROUND_HALF_UP) for ct, v in out.items()}
    log = [f"{ws.title}: {ct:<11}{cat:<48}{sub:<48}{float(v):>12,.2f}" for (ct, cat, sub), v in sorted(detail.items())]
    if external_lab:
        out["Prostheses"] = out.get("Prostheses", Decimal(0)) + external_lab
        log.append(f"{ws.title}: Prostheses  External LAB invoices (P&L line, not stock){'':<36}{float(external_lab):>12,.2f}")
    return dict(out), log


def implant_pools(report_path: str | Path) -> dict[str, Decimal]:
    """{pool: AED consumed} for the implant sub-categories (usage net of internal returns)."""
    from implant_codes import SUBCATEGORY_POOL
    wb = openpyxl.load_workbook(report_path, data_only=True, read_only=True)
    ws = next(s for s in wb.worksheets if s.title.lower().startswith("movement"))
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
    hi = {h: i for i, h in enumerate(hdr)}
    pools: dict[str, Decimal] = defaultdict(Decimal)
    for r in rows:
        if not r or r[hi["Transaction Type"]] not in CONSUMING:
            continue
        sub = str(r[hi["Sub-Category"]]).strip()
        for prefix, pool in SUBCATEGORY_POOL.items():
            if sub.startswith(prefix) and str(r[hi["Category"]]).startswith("05- Implants"):
                pools[pool] += -dec(r[hi["Transaction Value"]])
    return {k: v.quantize(D2, rounding=ROUND_HALF_UP) for k, v in pools.items()}


if __name__ == "__main__":
    p = sys.argv[1]
    ext = Decimal(sys.argv[2]) if len(sys.argv) > 2 else Decimal(0)
    split, log = cogs_split(p, ext)
    print("\n".join(log))
    print(f"\n{'Cost type':<12}{'AED':>14}")
    for ct, v in sorted(split.items(), key=lambda kv: -kv[1]):
        print(f"{ct:<12}{float(v):>14,.2f}")
    print(f"{'TOTAL':<12}{float(sum(split.values())):>14,.2f}")
