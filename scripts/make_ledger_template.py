#!/usr/bin/env python3
"""
Build the workbook Arjun fills in: FY2025 expenditure mapped to the sixteen DOH
cost types (Standard v2.1 Appendix 1), by month, with cost centres classified
direct or overhead (Stage 2). This is the cost-type anchor of the two-margin fit
(scripts/rebalance_two_view.py).

  python3 scripts/make_ledger_template.py --xmldir input/source/mysmyle_zip \
      --out input/Ledger_by_Cost_Type_FY2025_TEMPLATE.xlsx

Sheets: "Cost types by month" (to fill; each month must sum to the ledger
figure in scripts/gl_new.py), "Instructions", "Current implied" (what the sent
files imply per cost type, for reference only — not the ledger).
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from costsub_reader import read_submission  # noqa: E402
from costsub_schema import COST_TYPES  # noqa: E402
from gl_new import GL_2025, GL_2025_TOTAL  # noqa: E402

DESCRIPTIONS = {
    "SWNurs": "Salaries and wages — nursing (dental nurses)",
    "SWDoc": "Salaries and wages — medical (dentists, oral surgeons, orthodontists)",
    "SWAllied": "Salaries and wages — allied health incl. technicians (hygienists, radiographers, lab technicians)",
    "SWNonClin": "Salaries and wages — non-clinical staff recorded in DIRECT cost centres (chairside assistants, clinic clerks). Admin/management in overhead cost centres go to OHF.",
    "Laboratory": "Laboratory supplies and outsourced laboratory invoices (dental lab work)",
    "Imaging": "Imaging supplies (films, sensors, contrast) and outsourced imaging invoices",
    "Pharmacy": "Drugs, anaesthetics and pharmacy services",
    "Prostheses": "Implants, crowns, dentures and other implantable devices",
    "MS": "All other medical, surgical and dental supplies",
    "Hotel": "Cleaning products/services, linen and laundry, patient food",
    "GS": "All other goods and services recorded in DIRECT cost centres (e.g. rent, utilities, security if booked to the clinic rooms)",
    "Depreciation": "Depreciation and amortisation recorded in DIRECT cost centres (dental chairs, imaging units)",
    "OHF": "ALL expenses of OVERHEAD cost centres except depreciation: administration, reception/front office, management, HR, finance, IT, marketing, facility rent/utilities/maintenance booked to overhead centres",
    "OHDepreciation": "Depreciation and amortisation recorded in OVERHEAD cost centres",
    "OHC": "Overheads charged by a corporate office / shared services (group fees, head-office recharges)",
    "NABMs": "Non-activity based mandate funding by DOH, reported as negative cost (normally 0 for a private clinic)",
}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def implied_from_xml(xmldir: Path) -> dict:
    """{month: {cost type: AED}} as the sent files show them (reference only)."""
    out: dict = defaultdict(lambda: defaultdict(int))
    for p in sorted(xmldir.glob("*.xml")):
        m = re.search(r"_M(\d\d)_", p.name)
        if not m:
            continue
        month = int(m.group(1))
        s = read_submission(p)
        for c in s["claims"].values():
            for e in c["encounters"]:
                for t in COST_TYPES:
                    out[month][t] += e["money"].get(f"CostTypeDetails{t}", 0)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xmldir", default="input/source/mysmyle_zip")
    ap.add_argument("--out", default="input/Ledger_by_Cost_Type_FY2025_TEMPLATE.xlsx")
    args = ap.parse_args(argv)

    bold = Font(bold=True)
    fill_in = PatternFill("solid", fgColor="FFF2CC")
    head = PatternFill("solid", fgColor="D9E1F2")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cost types by month"
    ws.append(["Cost type", "Description (Standard v2.1 Appendix 1)", *MONTHS, "FY2025 total"])
    for c in ws[1]:
        c.font = bold
        c.fill = head
    for t in COST_TYPES:
        row = [t, DESCRIPTIONS[t]] + [None] * 12
        ws.append(row)
        r = ws.max_row
        ws.cell(r, 15).value = f"=SUM(C{r}:N{r})"
        for col in range(3, 15):
            ws.cell(r, col).fill = fill_in
            ws.cell(r, col).number_format = "#,##0.00"
        ws.cell(r, 15).number_format = "#,##0.00"
    first, last = 2, ws.max_row
    ws.append([])
    ws.append(["Total entered", ""] + [f"=SUM({openpyxl.utils.get_column_letter(c)}{first}:{openpyxl.utils.get_column_letter(c)}{last})" for c in range(3, 15)]
              + [f"=SUM(O{first}:O{last})"])
    tr = ws.max_row
    ws.append(["Ledger (Arjun, locked)", ""] + [float(GL_2025[m]) for m in range(1, 13)] + [float(GL_2025_TOTAL)])
    lr = ws.max_row
    ws.append(["Difference (must be 0.00)", ""] + [f"={openpyxl.utils.get_column_letter(c)}{tr}-{openpyxl.utils.get_column_letter(c)}{lr}" for c in range(3, 15)]
              + [f"=O{tr}-O{lr}"])
    ws.append([])
    ws.append(["Overhead ratio (OHF+OHDepreciation+OHC) / total", ""] +
              [f"=IFERROR(({openpyxl.utils.get_column_letter(c)}{first+12}+{openpyxl.utils.get_column_letter(c)}{first+13}+{openpyxl.utils.get_column_letter(c)}{first+14})/{openpyxl.utils.get_column_letter(c)}{tr},0)" for c in range(3, 15)]
              + [f"=IFERROR((O{first+12}+O{first+13}+O{first+14})/O{tr},0)"])
    for c in range(3, 16):
        ws.cell(ws.max_row, c).number_format = "0.0%"
    for r in (tr, lr, lr + 1, ws.max_row):
        for c in ws[r]:
            c.font = bold
        for c in range(3, 16):
            if r != ws.max_row:
                ws.cell(r, c).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 70
    for c in range(3, 16):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    ws.freeze_panes = "C2"

    ins = wb.create_sheet("Instructions")
    lines = [
        "Purpose: map every FY2025 expense line to one of the sixteen DOH cost types, by month, so the cost submission can be",
        "re-expressed with cost buckets and cost types that both equal the ledger (DOH SY02 rejection: Cost Bucket / Cost Type Match).",
        "",
        "1. Start from the same monthly ledger that gave the locked figures (row 'Ledger (Arjun, locked)'). Every month must reconcile to 0.00.",
        "2. First classify each cost centre as DIRECT (treats patients: dental surgeries, imaging room) or OVERHEAD (administration, reception/front office,",
        "   management, HR, finance, IT, marketing, facility services, head office).",
        "3. Expenses of OVERHEAD cost centres go to OHF (or OHDepreciation for their depreciation, OHC if charged by a corporate office).",
        "   Expenses of DIRECT cost centres go to the direct cost types by their nature (salaries by staff category, supplies, depreciation, GS).",
        "4. Rent, utilities, maintenance and security: if the premises are booked to an overhead cost centre, they are OHF; if booked to the clinic rooms, GS.",
        "5. Outsourced dental laboratory invoices are Laboratory; crowns/dentures/implants are Prostheses; anaesthetics and drugs are Pharmacy.",
        "6. Use accrual figures (Standard 3.1.3). Exclude balance-sheet items. Include third-party expenses not in the GL if relevant (3.1.2).",
        "7. The 'Current implied' sheet shows what the rejected files implied per cost type. It is for comparison only; do not copy it.",
        "8. DOH expects the overhead ratio between 20% and under 50%; below 20% draws a warning. Report what the ledger says.",
        "",
        "Return this file as input/Ledger_by_Cost_Type_FY2025.xlsx. The two-margin engine (scripts/rebalance_two_view.py) reads it.",
    ]
    for ln in lines:
        ins.append([ln])
    ins.column_dimensions["A"].width = 140

    cur = wb.create_sheet("Current implied")
    implied = implied_from_xml(Path(args.xmldir))
    cur.append(["Cost type", *MONTHS, "FY2025 total", "Share"])
    for c in cur[1]:
        c.font = bold
        c.fill = head
    grand = sum(sum(v.values()) for v in implied.values())
    for t in COST_TYPES:
        vals = [implied[m][t] / 1000 for m in range(1, 13)]
        cur.append([t, *vals, sum(vals), (sum(vals) * 1000 / grand) if grand else 0])
        for c in range(2, 15):
            cur.cell(cur.max_row, c).number_format = "#,##0.000"
        cur.cell(cur.max_row, 15).number_format = "0.00%"
    cur.append(["Total (cost-type view of the sent files)", *[sum(implied[m].values()) / 1000 for m in range(1, 13)], grand / 1000, 1])
    for c in cur[cur.max_row]:
        c.font = bold
    cur.append(["Note: the sent files carried only 64.5% of each claim in the cost-type view; the remaining 35.5% sat in the bucket view."])
    cur.column_dimensions["A"].width = 40
    for c in range(2, 16):
        cur.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
