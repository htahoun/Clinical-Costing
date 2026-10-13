#!/usr/bin/env python3
"""
Ledger-driven two-view engine (rulings of 10 Oct 2026; replaces the bucket-anchored fit).

  python3 scripts/rebuild_two_view.py --indir input/source/mysmyle_zip \
      --ledger input/Ledger_by_Cost_Type_FY2025.xlsx --outdir output/xml_resubmission \
      --activity input/source/Clinical_2025_one_sheet_Claude_Updated_Cherie.xlsx \
      --movement 1=input/source/Financial_Report_Jan_2025_1.xlsx [--zip]

Per month:
  1. Claim totals come from the sent file (scale-only costing, locked) and are
     tied to the ledger to the fils (largest claim absorbs the drift).
  2. Cost-type anchor: the ledger mapped to the sixteen cost types (Arjun).
  3. Prostheses are pre-assigned by billed units: implant stock on the implant
     codes (scripts/implant_codes.py); the rest of the Prostheses cost type
     (in-house CAD/CAM materials, external lab invoices) on the 48 prosthesis
     codes (scripts/prosthesis_codes.py); the dental technicians' share of
     SWAllied on the same claims, routed to the Prosthesis bucket. The clinic
     has no laboratory (ruling 13 Oct 2026): Laboratory is zero.
  4. A claim x cost-type matrix is fitted to both margins (claim totals, ledger
     cost types), seeded by the allocation table's own cells for each claim
     (what the code is made of), so a lab-heavy claim stays lab-heavy.
  5. Buckets are derived from the cost types by the fixed rules in
     scripts/bucket_rules.py, so buckets = cost types on every claim.
Exact integer-fils arithmetic for every written cell; identity fields untouched.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import openpyxl  # noqa: E402
from lxml import etree  # noqa: E402

from bucket_rules import OVERHEAD, rules_for_month  # noqa: E402
from costsub_reader import fils  # noqa: E402
from costsub_schema import BUCKETS, COST_TYPES, DATE_FMT, OVERHEAD_COST_TYPES, TZ_OFFSET_HOURS, fmt_money_fils  # noqa: E402
from gl_new import GL  # noqa: E402
from rebalance_two_view import MONTH_SHEETS, read_activity_units, read_ledger  # noqa: E402

# How a sent-file cell hints at cost types (seed weights). Bucket columns of the allocation table are
# reinterpreted by what they are made of; cost-type columns map to themselves.
SEED_FROM_BUCKET = {
    "Dental": {"MS": 1},                                   # supplies column
    "Laboratory": {"Prostheses": 1},                       # the sent OtherLab column = lab-made prostheses (ruling 13 Oct 2026)
    "Anaesthesia": {"Pharmacy": 0.5, "SWDoc": 0.5},
    "Imaging": {"Imaging": 0.4, "SWNurs": 0.3, "Depreciation": 0.3},
    "Prosthesis": {"Prostheses": 1},
    "OP": {"GS": 1},
    "Physician": {"SWDoc": 1},
    "Pharmacy": {"Pharmacy": 1},
    "Allied": {"SWAllied": 1},
    "Other": {"GS": 1},
}
EPS = 0.003          # share of the claim total added to every open cost type so the fit is always feasible
TOL = 1e-9


def read_lab_share(ledger_path: Path) -> dict[int, Decimal]:
    """Optional sheet 'Bucket shares': row 'SWAllied lab share' with 12 monthly values."""
    wb = openpyxl.load_workbook(ledger_path, data_only=True)
    if "Bucket shares" not in wb.sheetnames:
        return {}
    for r in wb["Bucket shares"].iter_rows(values_only=True):
        if r and r[0] == "SWAllied lab share":
            return {m: Decimal(str(r[m])) for m in range(1, 13) if r[m] is not None}
    return {}


def ipf(seed: list[dict[str, float]], rows: list[int], cols: dict[str, int], max_iter=2000):
    """Iterative proportional fitting of a claims x types matrix (floats) to integer-fils margins."""
    types = [t for t, v in cols.items() if v > 0]
    M = [{t: seed[i].get(t, 0.0) for t in types} for i in range(len(rows))]
    for i, r in enumerate(rows):
        if r > 0 and sum(M[i].values()) <= 0:
            raise SystemExit(f"claim {i} has no admissible cost type")
    for _ in range(max_iter):
        for i, r in enumerate(rows):
            s = sum(M[i].values())
            if s > 0:
                f = r / s
                for t in types:
                    M[i][t] *= f
        worst = 0.0
        for t in types:
            s = sum(M[i][t] for i in range(len(rows)))
            f = cols[t] / s if s > 0 else 0.0
            for i in range(len(rows)):
                M[i][t] *= f
            worst = max(worst, abs(f - 1))
        if worst < TOL:
            break
    else:
        raise SystemExit(f"fit did not converge (worst column factor off by {worst:.2e})")
    return M


def round_rows(M, rows):
    """Round each claim's cells to fils so they sum to the claim total (largest remainders)."""
    out = []
    for i, r in enumerate(rows):
        cells = {t: int(v) for t, v in M[i].items()}
        rem = {t: M[i][t] - cells[t] for t in cells}
        diff = r - sum(cells.values())
        for t in sorted(rem, key=lambda t: rem[t], reverse=True)[:max(diff, 0)]:
            cells[t] += 1
        out.append(cells)
    return out


def fix_columns(cells: list[dict[str, int]], cols: dict[str, int], log: list, fname: str):
    """Force every cost-type column to the ledger to the fils by moving fils on the richest claim."""
    resid = {t: cols[t] - sum(c.get(t, 0) for c in cells) for t in cols}
    if not any(resid.values()):
        return
    need = [t for t, v in resid.items() if v < 0]
    cand = [i for i, c in enumerate(cells) if all(c.get(t, 0) + resid[t] >= 0 for t in need)]
    i = max(cand, key=lambda i: sum(cells[i].values()))
    for t, v in resid.items():
        if v:
            cells[i][t] = cells[i].get(t, 0) + v
    log.append(f"{fname}: fils residuals {{{', '.join(f'{t}:{v:+d}' for t, v in resid.items() if v)}}} absorbed by claim index {i}")


def buckets_from_types(cells: dict[str, int], rules) -> tuple[dict[str, int], dict[str, int]]:
    """Exact-fils bucket cells for one claim from its cost-type cells; Σ buckets == Σ types."""
    direct = {b: Decimal(0) for b in BUCKETS}
    over = {b: Decimal(0) for b in BUCKETS}
    for t, v in cells.items():
        for b, share, comp in rules[t]:
            (over if comp == OVERHEAD else direct)[b] += Decimal(v) * share
    nd = {b: int(v.quantize(Decimal(1), rounding=ROUND_HALF_UP)) for b, v in direct.items()}
    no = {b: int(v.quantize(Decimal(1), rounding=ROUND_HALF_UP)) for b, v in over.items()}
    # keep the overhead component equal to the overhead cost types, and the claim equal in both views
    oh_types = sum(cells.get(t, 0) for t in OVERHEAD_COST_TYPES)
    no["Dental"] += oh_types - sum(no.values())
    nd["Dental"] += sum(cells.values()) - sum(nd.values()) - sum(no.values())
    return nd, no


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--indir", required=True)
    ap.add_argument("--ledger", required=True)
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--year", type=int, default=2025)
    ap.add_argument("--activity", help="activity master with treatment codes (monthly sheets)")
    ap.add_argument("--movement-book", help="full-year stock movement workbook (twelve 'Movement <Mon> 2025' sheets)")
    ap.add_argument("--movement", action="append", default=[], metavar="M=PATH")
    ap.add_argument("--zip", action="store_true")
    args = ap.parse_args(argv)
    movement = {m: args.movement_book for m in range(1, 13)} if args.movement_book else {}
    movement.update({int(s.split("=", 1)[0]): s.split("=", 1)[1] for s in args.movement})
    from cogs_from_movement import implant_pools
    from rebalance_two_view import preassign_prostheses

    ledger = read_ledger(Path(args.ledger))
    lab_share = read_lab_share(Path(args.ledger))
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone(timedelta(hours=TZ_OFFSET_HOURS))).strftime(DATE_FMT)
    log: list[str] = []
    print(f"{'File':<42}{'Claims':>7}{'Buckets':>16}{'Cost types':>16}{'Ledger':>16}{'Match':>7}{'OH%':>7}")
    written = []
    for f in sorted(Path(args.indir).glob("*.xml")):
        m = re.search(r"_M(\d\d)_", f.name)
        if not m:
            continue
        month = int(m.group(1))
        ledger_m = ledger[month]
        rules = rules_for_month(lab_share.get(month))
        tree = etree.parse(str(f))
        root = tree.getroot()
        encs, nodes = [], []
        for claim in root.findall("Claim"):
            cid = claim.findtext("ID")
            for enc in claim.findall("Encounter"):
                d = {b: fils(enc.find("CostBucketDirect").findtext(b)) for b in BUCKETS}
                o = {b: fils(enc.find("CostBucketOverheads").findtext(b)) for b in BUCKETS}
                t = {c: fils(enc.find("CostTypeDetails").findtext(c)) for c in COST_TYPES}
                encs.append({"cid": cid, "d": d, "o": o, "t": t, "total": sum(d.values()) + sum(o.values()) + sum(t.values())})
                nodes.append(enc)
        # 1. claim totals tied to the ledger
        gl = sum(ledger_m.values())
        drift = gl - sum(e["total"] for e in encs)
        if drift:
            i = max(range(len(encs)), key=lambda k: encs[k]["total"])
            encs[i]["total"] += drift
            log.append(f"{f.name}: claim totals {drift/1000:+.3f} AED vs ledger absorbed by claim {encs[i]['cid']}")
        # 2. seeds from the sent cells
        seeds = []
        open_types = [t for t, v in ledger_m.items() if v > 0]
        for e in encs:
            s: dict[str, float] = defaultdict(float)
            for t, v in e["t"].items():
                if v:
                    s[t] += v
            if s.get("OHF"):
                s["OHDepreciation"] += 0.05 * s["OHF"]
            for b, v in e["d"].items():
                if v:
                    for t, w in SEED_FROM_BUCKET.get(b, {"MS": 1}).items():
                        s[t] += v * w
            for t in open_types:
                s[t] += EPS * e["total"]
            seeds.append({t: s[t] for t in open_types})
        # 3. pre-assignment by billed units (ruling 13 Oct 2026: the clinic has no laboratory; every prosthesis
        #    code carries Prostheses). Pools: implant stock (fixtures, abutments, bone graft) on the implant codes;
        #    the rest of the month's Prostheses cost type (in-house CAD/CAM materials + external lab invoices) on
        #    the 48 prosthesis codes; the dental technicians' share of SWAllied on the same claims, routed to the
        #    Prosthesis bucket. Each pool is capped at the claim's remaining amount.
        pre: dict[int, int] = {}       # Prostheses fils per claim
        staff: dict[int, int] = {}     # SWAllied fils per claim that go to the Prosthesis bucket
        if args.activity:
            units = read_activity_units(args.activity, month)
            pools = implant_pools(movement[month], month if args.movement_book else None) if month in movement else {}
            implant_fils = sum(int((Decimal(v) * 1000).quantize(Decimal(1), rounding=ROUND_HALF_UP)) for v in pools.values())
            lab_fils = ledger_m.get("Prostheses", 0) - implant_fils
            if lab_fils < 0:
                raise SystemExit(f"{f.name}: implant stock {implant_fils/1000:,.2f} exceeds the Prostheses cost type {ledger_m.get('Prostheses', 0)/1000:,.2f}")
            if lab_fils:
                pools["lab"] = Decimal(lab_fils) / 1000
            log.append(f"{f.name}: Prostheses pools " + ", ".join(f"{k} {float(v):,.2f}" for k, v in pools.items())
                       + f"; coded claims in activity: {len(units)}")
            bucket_amt = [{"Dental": e["total"], "Prosthesis": 0} for e in encs]   # headroom = the whole claim
            pre = preassign_prostheses(encs, bucket_amt, units, pools, log, f.name)
            share = lab_share.get(month)
            if share and ledger_m.get("SWAllied", 0):
                staff_fils = int((Decimal(ledger_m["SWAllied"]) * share).quantize(Decimal(1), rounding=ROUND_HALF_UP))
                staff_units = {cid: {"dental technicians": u["lab"]} for cid, u in units.items() if u.get("lab")}
                staff = preassign_prostheses(encs, bucket_amt, staff_units, {"dental technicians": Decimal(staff_fils) / 1000}, log, f.name)
        pre_total, staff_total = sum(pre.values()), sum(staff.values())
        rows = [e["total"] - pre.get(i, 0) - staff.get(i, 0) for i, e in enumerate(encs)]
        cols = dict(ledger_m)
        if cols.get("Prostheses", 0) < pre_total:
            raise SystemExit(f"{f.name}: prosthesis pools exceed the Prostheses cost type")
        cols["Prostheses"] = cols.get("Prostheses", 0) - pre_total
        if cols.get("SWAllied", 0) < staff_total:
            raise SystemExit(f"{f.name}: dental technicians' pool exceeds the SWAllied cost type")
        cols["SWAllied"] = cols.get("SWAllied", 0) - staff_total
        cols = {t: v for t, v in cols.items() if v > 0}
        for s in seeds:
            for t in list(s):
                if t not in cols:
                    del s[t]
        # 4. fit and round
        M = ipf(seeds, rows, cols)
        cells = round_rows(M, rows)
        fix_columns(cells, cols, log, f.name)
        for i, p in pre.items():
            cells[i]["Prostheses"] = cells[i].get("Prostheses", 0) + p
        for i, p in staff.items():
            cells[i]["SWAllied"] = cells[i].get("SWAllied", 0) + p
        # 5. buckets from types; write
        B = T = OH = 0
        matrix: dict = defaultdict(lambda: defaultdict(Decimal))
        for i, (enc, c, e) in enumerate(zip(nodes, cells, encs)):
            full = {t: c.get(t, 0) for t in COST_TYPES}
            assert sum(full.values()) == e["total"], (e["cid"], sum(full.values()), e["total"])
            rem = dict(full)
            rem["SWAllied"] -= staff.get(i, 0)              # dental technicians' pay -> Prosthesis bucket, not the rule
            nd, no = buckets_from_types(rem, rules)
            nd["Prosthesis"] += staff.get(i, 0)
            matrix["SWAllied"]["Prosthesis"] += Decimal(staff.get(i, 0))
            assert sum(nd.values()) + sum(no.values()) == sum(full.values())
            for t, v in rem.items():
                for b, share, comp in rules[t]:
                    matrix[t][b] += Decimal(v) * share
            for name, vals in (("CostBucketDirect", nd), ("CostBucketOverheads", no), ("CostTypeDetails", full)):
                for child in enc.find(name):
                    child.text = fmt_money_fils(vals[child.tag])
            B += sum(nd.values()) + sum(no.values())
            T += sum(full.values())
            OH += sum(full[t] for t in OVERHEAD_COST_TYPES)
        td = root.find("Header/TransactionDate")
        if td is not None:
            td.text = now
        out = outdir / f.name
        tree.write(str(out), encoding="UTF-8", pretty_print=False)
        written.append(out)
        with open(outdir / f"matrix_M{month:02d}.csv", "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["bucket", "cost_type", "aed"])
            for t, row in matrix.items():
                for b, v in row.items():
                    if v:
                        w.writerow([b, t, f"{float(v) / 1000:.3f}"])
        bt = defaultdict(Decimal)
        for t, row in matrix.items():
            for b, v in row.items():
                bt[b] += v
        for b, v in bt.items():
            log.append(f"{f.name}: bucket {b} profile " + ", ".join(f"{t} {float(matrix[t][b] / v) * 100:.1f}%" for t in COST_TYPES if matrix[t].get(b)))
        ok = "OK" if B == T == gl else "FAIL"
        print(f"{f.name:<42}{len(encs):>7}{B/1000:>16,.3f}{T/1000:>16,.3f}{gl/1000:>16,.3f}{ok:>7}{OH/T*100:>6.1f}%")
    (outdir / "rebalance_log.txt").write_text("\n".join(log) + "\n")
    print(f"\n{len(written)} files written to {outdir}; log in rebalance_log.txt")
    if args.zip:
        zp = outdir.parent / f"{outdir.name}.zip"
        with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
            for w_ in written:
                z.write(w_, w_.name)
        print(f"zip: {zp}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
