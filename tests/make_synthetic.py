#!/usr/bin/env python3
"""
Synthetic activity master + allocation table so the pipeline can be proven
end-to-end without real patient data. Shapes mirror the real inputs (including
the locked-fix cases: 14611 duplicate at 378.00, identical 12701 rows, 72111
not summing to 100, orphan 12702, 99555, blank speciality, trailing-B serials).

  python3 tests/make_synthetic.py <outdir>
"""
import random
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

ALLOC_COLS = ["Anaesthesia", "Imaging", "Physician", "OP", "OtherLab", "Prosthesis", "Supplies",
              "SWNurs", "SWDoc", "SWAllied", "SWNonClin", "MS", "Hotel", "GS", "Depreciation", "OHF", "OHC"]
CODES = {"01901": 120.0, "02504": 95.5, "11101": 80.0, "12701": 150.0, "14611": 330.75, "29301": 210.0,
         "42311": 400.0, "27111": 260.0, "66301": 55.0, "72111": 180.0, "72119": 190.0, "97011": 70.0,
         "99555": 0.0}


def split(rnd):
    w = [rnd.random() for _ in ALLOC_COLS]
    w[6] *= 3; w[2] *= 4; w[15] *= 2; w[16] *= 2   # supplies, physician, overheads heavier
    s = sum(w)
    pct = [round(100 * x / s, 2) for x in w]
    pct[0] += round(100 - sum(pct), 2)
    return pct


def main(outdir):
    rnd = random.Random(7)
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Allocation"
    ws.append(["Code", "Description", "Cost/Unit", *ALLOC_COLS])
    for code, cost in CODES.items():
        pct = split(rnd)
        if code == "72111":
            pct[0] -= 7.5      # locked rescale case: sums to 92.5
        ws.append([code, f"desc {code}", cost, *pct])
    ws.append(["14611", "dup typo (really 14612)", 378.00, *split(rnd)])
    r = [c for c in ws.iter_rows(min_row=5, max_row=5, values_only=True)][0]
    ws.append(list(r))   # identical duplicate 12701
    wb.save(out / "allocation.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activity"
    ws.append(["Claim ID", "Serial", "FileNo.", "Doctor No", "Speciality", "StartDateTime", "EndDateTime",
               "Code", "Units", "Diagnosis", "TheatreAttendanceFlag", "Procedure", "Duration", "Type", "Payer", "Net", "Source"])
    docs = {"GD11525": "", "GD5633": "General Dentist", "GD9001": "Orthodontist", "GD9002": "Endodontist"}
    codes = list(CODES) + ["12702", "14612"]
    n = 0
    for m in range(1, 13):
        for k in range(rnd.randint(60, 90)):
            n += 1
            cid = f"{n}MF4159-{rnd.randrange(16**6):06x}"
            d = datetime(2025, m, rnd.randint(1, 28), rnd.randint(8, 20), rnd.choice([0, 15, 30, 45]))
            doc = rnd.choice(list(docs))
            first = f"{n * 10 + 1}{'B' if rnd.random() < 0.1 else ''}"
            for j in range(rnd.randint(1, 4)):
                code = rnd.choice(codes) if j else rnd.choice([c for c in codes if c != "99555"])
                ws.append([cid, first if j == 0 else f"{n * 10 + 1 + j}", f"{rnd.randint(1000, 9999)}", doc,
                           docs[doc] if rnd.random() > 0.2 else "", d.strftime("%d/%m/%Y %H:%M"),
                           (d + timedelta(minutes=30)).strftime("%d/%m/%Y %H:%M"), code, 1,
                           rnd.choice(["K02.9", "K04.0", "K05.1", "Z01.2"]), 0, "", 0, 1,
                           rnd.choice(["Daman", "ADNIC", "SelfPay"]), round(rnd.uniform(100, 900), 2), "Calibrate"])
    ws.append(["9999MF4159-zero00", "999991", "1234", "GD5633", "General Dentist", "15/06/2025 10:00",
               "15/06/2025 10:30", "99555", 1, "K02.9", 0, "", 0, 1, "SelfPay", 0, "test-zero"])
    wb.save(out / "activity.xlsx")
    print(f"wrote {out/'allocation.xlsx'} and {out/'activity.xlsx'} ({ws.max_row - 1} lines)")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "tests/synthetic")
